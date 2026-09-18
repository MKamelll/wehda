import io
from typing import Any

from django.contrib import admin
from django.db.models import Case, QuerySet, Value, When
from django.db.models import IntegerField
from django.http import FileResponse, HttpRequest, HttpResponse
from django.template.loader import render_to_string
from django.utils.http import content_disposition_header
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from weasyprint import HTML

from staff import models
from staff.forms import (
    DoctorDetailsFormSet,
    EmploymentGradeForm,
    EmploymentGradeFormSet,
    QualificationForm,
    QualificationFormSet,
    SecondedRecordFormSet,
)


@admin.action(description="PDF قوة الوحدة")
def workforce_pdf(
    model_admin: admin.ModelAdmin,
    request: HttpRequest,
    queryset: QuerySet[models.Employee],
) -> HttpResponse:

    unit_name = queryset[0].health_unit.name
    health_admin = queryset[0].health_unit.get_health_admin_display()
    headers = [
        "الاسم",
        "الوظيفة",
        "الدرجة",
        "الرقم القومي",
        "رقم التليفون",
        "ملاحظـــات",
    ]
    html_string = render_to_string(
        "staff/pdf_template.djhtml",
        context={
            "employees": queryset,
            "headers": headers,
            "unit_name": unit_name,
            "health_admin": health_admin,
            "empty_cols": range(1),
            "page_size": "A4 portrait",
        },
    )

    pdf_file = HTML(string=html_string).write_pdf()

    response = HttpResponse(pdf_file, content_type="application/pdf")
    response["Content-Disposition"] = (
        content_disposition_header(as_attachment=False, filename="قوة الوحدة.pdf") or ""
    )

    return response


@admin.action(description="PDF قوة الوحدة برقم الصرف")
def workforce_code_pdf(
    model_admin: admin.ModelAdmin,
    request: HttpRequest,
    queryset: QuerySet[models.Employee],
) -> HttpResponse:

    unit_name = queryset[0].health_unit.name
    health_admin = queryset[0].health_unit.get_health_admin_display()
    headers = [
        "الاسم",
        "الوظيفة",
        "الدرجة",
        "الرقم القومي",
        "رقم التليفون",
        "رقم الصرف",
        "ملاحظـــات",
    ]
    html_string = render_to_string(
        "staff/pdf_template.djhtml",
        context={
            "employees": queryset,
            "headers": headers,
            "unit_name": unit_name,
            "health_admin": health_admin,
            "empty_cols": range(2),
            "page_size": "A4 portrait",
        },
    )

    pdf_file = HTML(string=html_string).write_pdf()

    response = HttpResponse(pdf_file, content_type="application/pdf")
    response["Content-Disposition"] = (
        content_disposition_header(
            as_attachment=False, filename="قوة الوحدة برقم الصرف.pdf"
        )
        or ""
    )

    return response


@admin.action(description="xlsx العاملين بالوحدة")
def workforce_xlsx(
    model_admin: admin.ModelAdmin,
    request: HttpRequest,
    queryset: QuerySet[models.Employee],
) -> FileResponse:

    unit_name = queryset[0].health_unit.name
    health_admin = queryset[0].health_unit.get_health_admin_display()

    buffer = io.BytesIO()

    workbook = Workbook()
    worksheet = workbook.active

    worksheet.title = "العاملين"
    worksheet.views.sheetView[0].rightToLeft = True
    worksheet.append([f"الإدارة الصحية ب{health_admin}"])
    worksheet.append([f"قسم بيانات رعاية أساسية"])
    worksheet.append([f"وحدة {unit_name} الصحية"])
    worksheet.append([])
    worksheet.append([f"بيان العاملين بوحدة/ {unit_name}"])

    headers = [
        "الاسم",
        "أساسي/انتداب",
        "مكان العمل الأصلي للانتداب",
        "الدرجة",
        "الرقم القومي",
        "الوظيفة",
        "تخصص/للأطباء البشري",
        "سنة التخرج",
        "المؤهل",
        "جامعة تخرج",
        "رقم التليفون",
        "تاريخ الميلاد",
        "تاريخ استلام العمل داخل الوحدة",
        "تاريخ تعيين (أول تعيين)",
        "البريد الالكتروني",
        "العنوان (محل الإقامة)",
        "درجه علميه(دبلومه-ماجستير-دكتوراه-زماله)",
        "تاريخ الحصول عليه",
        "الجامعه",
        "تدريبات التي تم الحصول عليها(دورات)",
        "مكان الحصول علي الدورات التدريبيه",
        "رقم ترخيص مزاوله مهنه",
    ]
    worksheet.append(headers)

    worksheet.merge_cells(
        start_row=5, end_row=5, start_column=1, end_column=len(headers)
    )

    for i in queryset:
        current_secondment = i.current_secondment
        ow = str(current_secondment.original_workplace) if current_secondment else ""
        grade = i.current_grade.get_name_display().split(" ")[1]
        doctor = i.doctor_details.first()
        speciality = doctor.get_speciality_display() if doctor else ""
        qualification = i.qualification
        year = str(qualification.year or "") if qualification else ""
        institution = qualification.institution or ""
        address = i.address or ""
        additional_degrees = i.additional_degrees.select_related("institution")
        additional_degrees_names = "+".join([j.name for j in additional_degrees])
        additional_degrees_dates = "+".join(
            [str(j.date_of_certification) for j in additional_degrees]
        )
        additional_degrees_institutions = "+".join(
            [str(j.institution) for j in additional_degrees]
        )
        trainings = i.training_certifications.select_related("institution")
        trainings_names = "+".join([j.name for j in trainings])
        trainings_institutions = "+".join([str(j.institution or "") for j in trainings])
        row = [
            str(i.name),
            str(i.get_appointment_type_display()),
            ow,
            grade,
            str(i.national_id),
            str(i.get_job_title_display()),
            speciality,
            year,
            str(qualification),
            str(institution),
            str(i.phone_number),
            str(i.birth_date),
            str(i.job_start_date or ""),
            str(i.first_hire_date or ""),
            str(i.email or ""),
            str(address),
            additional_degrees_names,
            additional_degrees_dates,
            additional_degrees_institutions,
            trainings_names,
            trainings_institutions,
            i.practice_license_number,
        ]
        worksheet.append(row)

    font = Font(name="Arial", size=12, bold=True, italic=False, color="000000")
    grey_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    yellow_fill = PatternFill(
        start_color="FFFF00", end_color="FFFF00", fill_type="solid"
    )
    border_side = Side(border_style="thin", color="000000")
    full_border = Border(
        left=border_side, right=border_side, top=border_side, bottom=border_side
    )

    for row in worksheet:
        for cell in row:
            cell.font = font
            if cell.row == 5:
                cell.fill = yellow_fill
            if cell.row >= 5 and cell.column <= len(headers):
                cell.border = full_border
            if cell.row == 6:
                cell.fill = grey_fill
                cell.alignment = Alignment(
                    vertical="center", horizontal="center", wrap_text=True
                )
            else:
                cell.alignment = Alignment(vertical="center", horizontal="center")

    workbook.save(buffer)
    buffer.seek(0)

    return FileResponse(buffer, as_attachment=True, filename="العاملين.xlsx")


@admin.action(description="PDF العاملين بالوحدة")
def workforce_full_pdf(
    model_admin: admin.ModelAdmin,
    request: HttpRequest,
    queryset: QuerySet[models.Employee],
) -> HttpResponse:

    unit_name = queryset[0].health_unit.name
    health_admin = queryset[0].health_unit.get_health_admin_display()
    tables: list[dict[str, list[str] | list[list[str]]]] = [
        {
            "headers": [
                "الاسم",
                "أساسي/انتداب",
                "مكان العمل الأصلي للانتداب",
                "الدرجة",
                "الرقم القومي",
            ],
            "rows": [],
        },
        {
            "headers": [
                "الاسم",
                "الوظيفة",
                "تخصص/للأطباء البشري",
                "سنة التخرج",
                "جامعة تخرج",
                "رقم التليفون",
            ],
            "rows": [],
        },
        {
            "headers": [
                "الاسم",
                "تاريخ الميلاد",
                "تاريخ استلام العمل داخل الوحدة",
                "تاريخ تعيين (أول تعيين)",
                "البريد الالكتروني",
            ],
            "rows": [],
        },
        {
            "headers": [
                "الاسم",
                "العنوان (محل الإقامة)",
                "درجه علميه(دبلومه-ماجستير-دكتوراه-زماله)",
                "تاريخ الحصول عليه",
                "الجامعه",
            ],
            "rows": [],
        },
        {
            "headers": [
                "الاسم",
                "تدريبات التي تم الحصول عليها(دورات)",
                "مكان الحصول علي الدورات التدريبيه",
                "رقم ترخيص مزاوله مهنه",
            ],
            "rows": [],
        },
    ]

    table_1_rows = []
    for i in queryset:
        row = []
        row.append(str(i.name))
        row.append(i.get_appointment_type_display())
        current_secondment = i.current_secondment
        ow = (
            str(current_secondment.original_workplace or "")
            if current_secondment
            else ""
        )
        row.append(ow)
        grade = i.current_grade.get_name_display().split(" ")[1]
        row.append(grade)
        row.append(str(i.national_id))
        table_1_rows.append(row)

    tables[0]["rows"] = table_1_rows

    table_2_rows = []
    for i in queryset:
        row = []
        row.append(str(i.name))
        row.append(i.get_job_title_display())
        doctor = i.doctor_details.first()
        speciality = doctor.get_speciality_display() if doctor else ""
        row.append(speciality)
        qualification = i.qualification
        year = str(qualification.year or "") if qualification else ""
        row.append(year)
        institution = qualification.institution or ""
        row.append(str(institution))
        row.append(str(i.phone_number))
        table_2_rows.append(row)

    tables[1]["rows"] = table_2_rows

    table_3_rows = []
    for i in queryset:
        row = []
        row.append(str(i.name))
        row.append(str(i.birth_date))
        row.append(str(i.job_start_date or ""))
        row.append(str(i.first_hire_date or ""))
        row.append(i.email or "")
        table_3_rows.append(row)

    tables[2]["rows"] = table_3_rows

    table_4_rows = []
    for i in queryset:
        row = []
        row.append(str(i.name))
        address = i.address or ""
        row.append(str(address))
        additional_degrees = i.additional_degrees.select_related("institution")
        additional_degrees_names = "+".join([j.name for j in additional_degrees])
        additional_degrees_dates = "+".join(
            [str(j.date_of_certification) for j in additional_degrees]
        )
        additional_degrees_institutions = "+".join(
            [str(j.institution) for j in additional_degrees]
        )
        row.append(additional_degrees_names)
        row.append(additional_degrees_dates)
        row.append(additional_degrees_institutions)
        table_4_rows.append(row)

    tables[3]["rows"] = table_4_rows

    table_5_rows = []
    for i in queryset:
        row = []
        row.append(str(i.name))
        trainings = i.training_certifications.select_related("institution")
        trainings_names = "+".join([j.name for j in trainings])
        trainings_institutions = "+".join([str(j.institution or "") for j in trainings])
        row.append(trainings_names)
        row.append(trainings_institutions)
        row.append(i.practice_license_number or "")
        table_5_rows.append(row)

    tables[4]["rows"] = table_5_rows

    html_string = render_to_string(
        "staff/pdf_full_template.djhtml",
        context={
            "tables": tables,
            "unit_name": unit_name,
            "health_admin": health_admin,
            "empty_cols": range(0),
            "page_size": "A4 landscape",
        },
    )

    pdf_file = HTML(string=html_string).write_pdf()

    response = HttpResponse(pdf_file, content_type="application/pdf")
    response["Content-Disposition"] = (
        content_disposition_header(as_attachment=False, filename="العاملين.pdf") or ""
    )

    return response


@admin.register(models.Institution)
class InstituitionAdmin(admin.ModelAdmin[models.Institution]):
    search_fields = ("name",)

    def has_module_permission(self, _request: HttpRequest) -> bool:
        return False


@admin.register(models.Address)
class AddressAdmin(admin.ModelAdmin[models.Address]):
    search_fields = ("street", "village", "city", "governate")

    def has_module_permission(self, _request: HttpRequest) -> bool:
        return False


@admin.register(models.HealthUnit)
class HealthUnitAdmin(admin.ModelAdmin[models.HealthUnit]):
    list_display = (
        "name",
        "code",
        "health_admin",
        "number_of_residents",
        "number_of_clinics",
    )

    search_fields = ("name", "code", "health_admin")
    autocomplete_fields = ("address",)


class DoctorDetailsInline(admin.StackedInline[models.DoctorDetails, models.Employee]):
    model = models.DoctorDetails
    formset = DoctorDetailsFormSet
    can_delete = True
    extra = 0
    max_num = 1


class QualificationInline(admin.StackedInline[models.Qualification, models.Employee]):
    model = models.Qualification
    formset = QualificationFormSet
    form = QualificationForm
    can_delete = True
    extra = 0
    autocomplete_fields = ["institution"]


class EmploymentGradeInline(
    admin.StackedInline[models.EmploymentGrade, models.Employee]
):
    model = models.EmploymentGrade
    formset = EmploymentGradeFormSet
    form = EmploymentGradeForm
    can_delete = True
    extra = 0


class AddonTaskInline(admin.StackedInline[models.AddonTask, models.Employee]):
    model = models.AddonTask
    can_delete = True
    extra = 0


class SecondedRecordInline(admin.StackedInline[models.SecondedRecord, models.Employee]):
    model = models.SecondedRecord
    formset = SecondedRecordFormSet
    ordering = ("-start_date",)
    can_delete = False
    extra = 0


class AdditionalDegreeInline(
    admin.StackedInline[models.AdditionalDegree, models.Employee]
):
    model = models.AdditionalDegree
    extra = 0
    can_delete = True
    autocomplete_fields = ["institution"]


class TrainingCertificationInline(
    admin.StackedInline[models.TrainingCertification, models.Employee]
):
    model = models.TrainingCertification
    extra = 0
    can_delete = True
    autocomplete_fields = ["institution"]


@admin.register(models.Employee)
class EmployeeModelAdmin(admin.ModelAdmin[models.Employee]):
    list_display = (
        "name",
        "sorted_job_title",
        "appointment_type",
    )

    list_filter = ("appointment_type", "job_title", "is_active")

    search_fields = ("name", "national_id", "addon_tasks__name")
    autocomplete_fields = ("address",)
    inlines = [
        QualificationInline,
        EmploymentGradeInline,
        DoctorDetailsInline,
        SecondedRecordInline,
        AdditionalDegreeInline,
        TrainingCertificationInline,
        AddonTaskInline,
    ]
    actions = [workforce_pdf, workforce_code_pdf, workforce_xlsx, workforce_full_pdf]
    ordering = ["-id"]

    def get_queryset(self, request):
        queryset = super().get_queryset(request)

        queryset = queryset.annotate(
            job_group_weight=Case(
                When(
                    job_title=models.Employee.JobTitlesChoices.HEAD_OF_UNIT,
                    then=Value(1),
                ),
                When(job_title=models.Employee.JobTitlesChoices.DOCTOR, then=Value(2)),
                When(
                    job_title=models.Employee.JobTitlesChoices.HEAD_DENTIST,
                    then=Value(3),
                ),
                When(job_title=models.Employee.JobTitlesChoices.DENTIST, then=Value(4)),
                When(
                    job_title=models.Employee.JobTitlesChoices.HEAD_PHARMACIST,
                    then=Value(5),
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.PHARMACIST, then=Value(6)
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.HEAD_NURSE, then=Value(7)
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.NURSE_CONSULTANT,
                    then=Value(8),
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.SENIOR_NURSING_SPECIALIST,
                    then=Value(9),
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.NURSING_SPECIALIST,
                    then=Value(10),
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.NURSING_TECHNICIAN,
                    then=Value(11),
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.HEALTH_INSPECTOR,
                    then=Value(12),
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.LAB_TECHNICIAN,
                    then=Value(13),
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.HEALTH_REGISTER,
                    then=Value(14),
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.GENERAL_REGISTER,
                    then=Value(15),
                ),
                When(
                    job_title=models.Employee.JobTitlesChoices.COMMUNITY_HEALTH_WORKER,
                    then=Value(16),
                ),
                default=Value(99),
                output_field=IntegerField(),
            )
        )
        return queryset

    @admin.display(ordering="job_group_weight", description="الوظيفة")
    def sorted_job_title(self, obj: models.Employee):
        return obj.get_job_title_display()
